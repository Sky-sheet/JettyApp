# -*- coding: utf-8 -*-
from kivymd.uix.screen import MDScreen
from kivymd.uix.boxlayout import MDBoxLayout
from kivymd.uix.label import MDLabel
from kivymd.uix.textfield import MDTextField
from kivymd.uix.button import MDButton, MDButtonText
from kivy.app import App
from utils.helpers import current_jalali_year, current_jalali_month, get_data_dir
from utils.rtl import rtl

F = "Roboto"


def field_label(text):
    return MDLabel(
        text=rtl(text), halign="right", size_hint_y=None, height="26dp",
        font_name=F, theme_text_color="Secondary",
    )


class ExportScreen(MDScreen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.name = "export"

        root = MDBoxLayout(orientation="vertical", padding="16dp", spacing="10dp")
        root.add_widget(MDButton(
            MDButtonText(text=rtl("بازگشت"), font_name=F), style="text", size_hint_x=1,
            on_release=lambda x: setattr(self.manager, "current", "home"),
        ))
        root.add_widget(MDLabel(text=rtl("خروجی گزارش Excel"), halign="center",
                                size_hint_y=None, height="40dp", font_name=F))

        root.add_widget(field_label("سال"))
        self.year_f = MDTextField(text=str(current_jalali_year()), font_name=F)
        root.add_widget(self.year_f)

        root.add_widget(field_label("ماه (01 تا 12)"))
        self.month_f = MDTextField(text=f"{current_jalali_month():02d}", font_name=F)
        root.add_widget(self.month_f)

        root.add_widget(MDButton(
            MDButtonText(text=rtl("تولید فایل Excel"), font_name=F), style="filled", size_hint_x=1,
            on_release=self._export,
        ))
        self.status = MDLabel(text="", size_hint_y=None, height="80dp", halign="center", font_name=F)
        root.add_widget(self.status)
        root.add_widget(MDLabel(
            text=rtl("فایل در پوشه داده های برنامه ذخیره می شود"),
            size_hint_y=None, height="40dp", halign="center",
            theme_text_color="Secondary", font_name=F,
        ))
        self.add_widget(root)

    def _export(self, *args):
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment
        except ImportError:
            self.status.text = rtl("کتابخانه openpyxl نصب نیست\npip install openpyxl")
            return

        year = (self.year_f.text or "").strip()
        month = (self.month_f.text or "").strip()
        if not year or not month:
            self.status.text = rtl("سال و ماه را وارد کنید")
            return

        app = App.get_running_app()
        rows = app.db.get_employees_by_month(year, month)
        if not rows:
            self.status.text = rtl("داده ای برای این ماه وجود ندارد")
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "گزارش"
        try:
            ws.sheet_view.rightToLeft = True
        except Exception:
            pass

        headers = ["ردیف", "نام", "شماره پرسنلی", "استخدام", "تاریخ", "اسکله"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for i, row in enumerate(rows, 1):
            vals = [i, row[1], row[2], row[3], row[4], row[5]]
            for col, v in enumerate(vals, 1):
                cell = ws.cell(row=i + 1, column=col, value=v)
                cell.alignment = Alignment(horizontal="center")

        data_dir = get_data_dir(app)
        path = data_dir / f"report_{year}_{month}.xlsx"
        wb.save(str(path))
        self.status.text = rtl(f"ذخیره شد:\n{path}")
